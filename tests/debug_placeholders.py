from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

# 👉 Change this to your template path
template_path = r"C:\Users\sesa747199\OneDrive - Schneider Electric\1.0 Technical Support\1. Product_Category_with_Case\0. PPT Template\Case_Sample_PPT_Node.pptx"

prs = Presentation(template_path)
slide = prs.slides[0]

print("===== PLACEHOLDERS ON COVER SLIDE =====")
for idx, shape in enumerate(slide.shapes):
    if shape.is_placeholder:
        ph = shape.placeholder_format
        print(
            f"Shape {idx}: "
            f"type={ph.type}, "
            f"idx={ph.idx}, "
            f"name={shape.name!r}, "
            f"text={shape.text!r}"
        )
print("========================================")
